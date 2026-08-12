#!/usr/bin/env python3
"""Manually publish a historical AMDC Excel workbook into its dated Feishu sheet."""
import argparse
import contextlib
import csv
import datetime as dt
import errno
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


CHUNK_ROWS = 25
CATEGORY_COLORS = {
    '超休闲': ('#FFD180', '#5D3600'),
    '休闲': ('#A5D6A7', '#174D25'),
    '壁纸': ('#CE93D8', '#4A1D54'),
    'Launcher': ('#90CAF9', '#0D3B66'),
    '杀毒软件、清理': ('#FFAB91', '#6D1B0D'),
    '文件恢复': ('#FFF59D', '#5F4B00'),
    'PDF阅读器': ('#80CBC4', '#004D40'),
}

REPO_ROOT = Path(__file__).resolve().parents[3]
PROJECT_DIR = Path(__file__).resolve().parents[1]


def configured_feishu_url():
    value = os.environ.get('AMDC_FEISHU_SHEET_URL', '').strip()
    if value:
        return value
    config_path = Path(os.environ.get('AMDC_CONFIG_FILE') or PROJECT_DIR / 'amdc-config.json')
    if not config_path.is_file():
        return ''
    try:
        config = json.loads(config_path.read_text(encoding='utf-8-sig'))
    except (OSError, json.JSONDecodeError):
        return ''
    integrations = config.get('integrations') if isinstance(config, dict) else None
    if not isinstance(integrations, dict):
        return ''
    return str(integrations.get('feishuSheetUrl') or '').strip()


def fail(message):
    print(json.dumps({'ok': False, 'error': message}, ensure_ascii=True))
    raise SystemExit(1)


def progress(percent, message):
    print(json.dumps({'event': 'progress', 'percent': percent, 'message': message}, ensure_ascii=True), flush=True)


def cli(args, *, input_data=None, allow_failure=False):
    env = os.environ.copy()
    env['LARKSUITE_CLI_NO_UPDATE_NOTIFIER'] = '1'
    env['LARKSUITE_CLI_NO_SKILLS_NOTIFIER'] = '1'
    configured_bin = os.environ.get('LARK_CLI_BIN')
    local_bin = REPO_ROOT / 'node_modules' / '.bin' / ('lark-cli.cmd' if os.name == 'nt' else 'lark-cli')
    cli_bin = configured_bin or (str(local_bin) if local_bin.exists() else None) or shutil.which('lark-cli.cmd' if os.name == 'nt' else 'lark-cli')
    if not cli_bin:
        raise RuntimeError('lark-cli is not installed. Run npm install from the AMTools root.')
    result = subprocess.run(
        [cli_bin, 'sheets', *args], input=input_data, text=True,
        capture_output=True, encoding='utf-8', errors='replace', env=env,
    )
    if result.returncode and not allow_failure:
        detail = (result.stderr or result.stdout or 'lark-cli failed').strip()
        raise RuntimeError(detail[-2000:])
    try:
        data = json.loads(result.stdout or '{}')
    except json.JSONDecodeError:
        data = {'ok': False, 'raw': result.stdout[-2000:]}
    if not data.get('ok', result.returncode == 0) and not allow_failure:
        raise RuntimeError(json.dumps(data, ensure_ascii=False)[-2000:])
    return data


def color_hex(color):
    if not color or color.type != 'rgb' or not color.rgb:
        return ''
    value = color.rgb[-6:]
    return f'#{value.upper()}' if value and value != '000000' else ''


def border_side(side):
    if not side or not side.style:
        return None
    style_map = {
        'thin': ('solid', 'thin'), 'medium': ('solid', 'medium'), 'thick': ('solid', 'thick'),
        'double': ('double', 'medium'), 'dashed': ('dashed', 'thin'), 'dotted': ('dotted', 'thin'),
    }
    style, weight = style_map.get(side.style, ('solid', 'thin'))
    result = {'style': style, 'weight': weight}
    color = color_hex(side.color)
    if color:
        result['color'] = color
    return result


def report_font_size(row_index):
    """Return the desired report font size in Excel points."""
    if row_index == 1:
        return 14
    if row_index in (2, 3):
        return 9
    if row_index >= 5:
        return 11
    return None


def points_to_feishu_pixels(points):
    """Map the requested Feishu UI font sizes to the Sheets API pixel values."""
    preset = {14: 18, 9: 12, 11: 14}
    if points in preset:
        return preset[points]
    return max(1, round(points * 4 / 3))


def normalized_hyperlink(cell):
    target = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else ''
    match = re.match(r'^(https://apps\.apple\.com)/app/(id\d+)(.*)$', target)
    if not match:
        return target
    headquarters = str(cell.parent.cell(cell.row, 17).value or '').strip().lower()
    storefront = headquarters if re.fullmatch(r'[a-z]{2}', headquarters) else 'us'
    return f'{match.group(1)}/{storefront}/app/{match.group(2)}{match.group(3)}'


def cell_payload(cell, no_wrap=False):
    payload = {}
    value = cell.value
    hyperlink = normalized_hyperlink(cell)
    if hyperlink and value is not None:
        payload['rich_text'] = [{'type': 'link', 'text': str(value), 'link': hyperlink}]
    elif value is not None:
        if cell.data_type == 'f':
            payload['formula'] = str(value if str(value).startswith('=') else '=' + str(value))
        elif isinstance(value, (dt.datetime, dt.date, dt.time)):
            payload['value'] = value.isoformat()
        elif isinstance(value, (str, int, float, bool)):
            payload['value'] = value
        else:
            payload['value'] = str(value)

    styles = {}
    category_colors = CATEGORY_COLORS.get(str(cell.value or '')) if cell.column == 3 and cell.row >= 6 else None
    background = color_hex(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else ''
    if category_colors:
        styles['background_color'] = category_colors[0]
    elif background:
        styles['background_color'] = background
    font_color = color_hex(cell.font.color)
    if category_colors:
        styles['font_color'] = category_colors[1]
    elif font_color:
        styles['font_color'] = font_color
    fixed_font_size = report_font_size(cell.row)
    if fixed_font_size:
        styles['font_size'] = points_to_feishu_pixels(fixed_font_size)
    elif cell.font.sz:
        styles['font_size'] = points_to_feishu_pixels(float(cell.font.sz))
    if cell.font.bold:
        styles['font_weight'] = 'bold'
    if cell.font.italic:
        styles['font_style'] = 'italic'
    if cell.font.underline:
        styles['font_line'] = 'underline'
    align_h = {'center': 'center', 'left': 'left', 'right': 'right'}.get(cell.alignment.horizontal)
    align_v = {'center': 'middle', 'top': 'top', 'bottom': 'bottom'}.get(cell.alignment.vertical)
    if align_h:
        styles['horizontal_alignment'] = align_h
    if align_v:
        styles['vertical_alignment'] = align_v
    if no_wrap:
        styles['word_wrap'] = 'overflow'
    elif cell.alignment.wrap_text:
        styles['word_wrap'] = 'auto-wrap'
    if cell.number_format and cell.number_format != 'General':
        styles['number_format'] = cell.number_format
    if styles:
        payload['cell_styles'] = styles

    borders = {name: border_side(getattr(cell.border, name)) for name in ('top', 'bottom', 'left', 'right')}
    borders = {name: side for name, side in borders.items() if side}
    if borders:
        payload['border_styles'] = borders
    return payload


def display_width(value):
    """Approximate rendered width: CJK/full-width characters take two slots."""
    if value is None:
        return 0
    lines = str(value).splitlines() or ['']
    return max(sum(2 if unicodedata.east_asian_width(char) in ('F', 'W', 'A') else 1 for char in line) for line in lines)


def column_pixels(ws, col_index):
    """Fit a table column to its longest visible cell, without wrapping its content."""
    content_width = max(display_width(ws.cell(row, col_index).value) for row in range(5, ws.max_row + 1))
    fitted_pixels = content_width * 8 + 20
    return max(48, fitted_pixels)


def normalized_compare_value(value):
    """Normalize the two columns used to detect an unchanged report."""
    if value is None:
        return ''
    text = ' '.join(str(value).replace('\r\n', '\n').replace('\r', '\n').split())
    if re.fullmatch(r'-?\d+\.0+', text):
        return text.split('.', 1)[0]
    return text


def column_values_from_csv(raw):
    values = []
    for row in csv.reader(io.StringIO(raw or ''), skipinitialspace=True):
        values.append(normalized_compare_value(row[0] if row else ''))
    while values and values[-1] == '':
        values.pop()
    return values


def report_columns_match(ws, max_row, url_args, sheet_id, existing_rows):
    """Return True only when title and current-rank columns match in order."""
    if existing_rows < 5 or max_row < 5:
        return False
    local = {
        2: [normalized_compare_value(ws.cell(row, 2).value) for row in range(5, max_row + 1)],
        5: [normalized_compare_value(ws.cell(row, 5).value) for row in range(5, max_row + 1)],
    }
    remote_end = max(existing_rows, max_row)
    for col in (2, 5):
        try:
            data = cli([
                '+csv-get', *url_args, '--sheet-id', sheet_id,
                '--range', f'{get_column_letter(col)}5:{get_column_letter(col)}{remote_end}',
                '--include-row-prefix=false', '--max-chars', '500000',
            ])['data']
            if data.get('has_more'):
                return False
            remote = column_values_from_csv(data.get('annotated_csv', ''))
        except Exception:
            return False
        if remote != local[col]:
            return False
    return True


SHEET_ORDER_LOCK_TIMEOUT_SECONDS = 300
SHEET_ORDER_LOCK_RETRY_SECONDS = 0.25


@contextlib.contextmanager
def sheet_order_lock(url, timeout_seconds=SHEET_ORDER_LOCK_TIMEOUT_SECONDS):
    """Serialize workbook tab ordering across parallel sync processes."""
    lock_name = hashlib.sha256(url.encode('utf-8')).hexdigest()[:20]
    lock_path = Path(tempfile.gettempdir()) / f'amdc-feishu-sheet-order-{lock_name}.lock'
    with lock_path.open('a+b') as lock_file:
        lock_file.seek(0, os.SEEK_END)
        if lock_file.tell() == 0:
            lock_file.write(b'0')
            lock_file.flush()
        lock_file.seek(0)
        if os.name == 'nt':
            import msvcrt
            deadline = time.monotonic() + max(0, float(timeout_seconds))
            retryable_errors = {errno.EACCES, errno.EAGAIN, errno.EDEADLK}
            while True:
                try:
                    msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError as error:
                    if error.errno not in retryable_errors:
                        raise
                    if time.monotonic() >= deadline:
                        raise TimeoutError('Timed out waiting for the Feishu workbook sheet-order lock') from error
                    time.sleep(SHEET_ORDER_LOCK_RETRY_SECONDS)
            try:
                yield
            finally:
                lock_file.seek(0)
                msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def sheet_title(sheet):
    return str(sheet.get('title') or sheet.get('sheet_name') or '')


def desired_sheet_order(sheets):
    """Put dated sheets first, newest to oldest, preserving other sheet order."""
    dated = [sheet for sheet in sheets if re.fullmatch(r'\d{8}', sheet_title(sheet))]
    other = [sheet for sheet in sheets if not re.fullmatch(r'\d{8}', sheet_title(sheet))]
    return sorted(dated, key=sheet_title, reverse=True) + other


def sort_date_sheets_desc(url_args, url):
    """Move workbook tabs into deterministic date-descending order and verify it."""
    with sheet_order_lock(url):
        current = list(cli(['+workbook-info', *url_args])['data'].get('sheets', []))
        expected = desired_sheet_order(current)
        for target_index, wanted in enumerate(expected):
            current_index = next(
                index for index, sheet in enumerate(current)
                if sheet.get('sheet_id') == wanted.get('sheet_id')
            )
            if current_index == target_index:
                continue
            cli([
                '+sheet-move', *url_args,
                '--sheet-id', wanted['sheet_id'],
                '--source-index', str(current_index),
                '--index', str(target_index),
            ])
            moved = current.pop(current_index)
            current.insert(target_index, moved)

        actual = list(cli(['+workbook-info', *url_args])['data'].get('sheets', []))
        actual_titles = [sheet_title(sheet) for sheet in actual]
        expected_titles = [sheet_title(sheet) for sheet in expected]
        if actual_titles != expected_titles:
            raise RuntimeError(
                'Feishu sheet order validation failed: '
                f'expected {expected_titles}, got {actual_titles}'
            )
        return actual_titles


def sort_workbook_date_sheets(url_args):
    """Sort the complete workbook using the URL already passed to lark-cli."""
    try:
        url = url_args[url_args.index('--url') + 1]
    except (ValueError, IndexError) as error:
        raise RuntimeError('Feishu workbook URL is missing from sheet operation arguments') from error
    return sort_date_sheets_desc(url_args, url)


def transaction_sheets(url_args):
    return list(cli(['+workbook-info', *url_args])['data'].get('sheets', []))


def transaction_sheet(sheets, title):
    return next((sheet for sheet in sheets if sheet_title(sheet) == title), None)


def transaction_result(action, target_name, backup_name, **extra):
    payload = {
        'ok': True,
        'transaction': action,
        'sheetName': target_name,
        'backupName': backup_name,
    }
    payload.update(extra)
    print(json.dumps(payload, ensure_ascii=True))


def prepare_transaction(url_args, target_name, backup_name):
    sheets = transaction_sheets(url_args)
    target = transaction_sheet(sheets, target_name)
    existing_backup = transaction_sheet(sheets, backup_name)
    if existing_backup:
        raise RuntimeError(f'transaction backup already exists: {backup_name}')
    if not target:
        transaction_result('prepare', target_name, backup_name, existed=False, originalIndex=-1)
        return
    original_index = next(index for index, sheet in enumerate(sheets) if sheet.get('sheet_id') == target.get('sheet_id'))
    cli([
        '+sheet-copy', *url_args, '--sheet-id', target['sheet_id'],
        '--title', backup_name, '--index', str(original_index),
    ])
    copied = transaction_sheet(transaction_sheets(url_args), backup_name)
    if not copied:
        raise RuntimeError(f'transaction backup could not be verified: {backup_name}')
    transaction_result('prepare', target_name, backup_name, existed=True, originalIndex=original_index)


def rollback_transaction(url_args, target_name, backup_name, existed, original_index):
    sheets = transaction_sheets(url_args)
    target = transaction_sheet(sheets, target_name)
    backup = transaction_sheet(sheets, backup_name)
    if backup:
        if target and target.get('sheet_id') != backup.get('sheet_id'):
            cli(['+sheet-delete', *url_args, '--sheet-id', target['sheet_id'], '--yes'])
        cli(['+sheet-rename', *url_args, '--sheet-id', backup['sheet_id'], '--title', target_name])
        if original_index is not None and int(original_index) >= 0:
            current = transaction_sheets(url_args)
            restored = transaction_sheet(current, target_name)
            current_index = next((index for index, sheet in enumerate(current) if sheet.get('sheet_id') == restored.get('sheet_id')), -1)
            if restored and current_index >= 0 and current_index != int(original_index):
                cli([
                    '+sheet-move', *url_args, '--sheet-id', restored['sheet_id'],
                    '--source-index', str(current_index), '--index', str(int(original_index)),
                ])
        sheet_order = sort_workbook_date_sheets(url_args)
        transaction_result('rollback', target_name, backup_name, restored=True, sheetOrder=sheet_order)
        return
    if not existed and target:
        cli(['+sheet-delete', *url_args, '--sheet-id', target['sheet_id'], '--yes'])
        sheet_order = sort_workbook_date_sheets(url_args)
        transaction_result('rollback', target_name, backup_name, restored=True, deletedCreatedSheet=True, sheetOrder=sheet_order)
        return
    if existed:
        raise RuntimeError(f'transaction backup missing: {backup_name}')
    sheet_order = sort_workbook_date_sheets(url_args)
    transaction_result('rollback', target_name, backup_name, restored=True, unchanged=True, sheetOrder=sheet_order)


def cleanup_transaction(url_args, backup_name):
    backup = transaction_sheet(transaction_sheets(url_args), backup_name)
    if backup:
        cli(['+sheet-delete', *url_args, '--sheet-id', backup['sheet_id'], '--yes'])
    sheet_order = sort_workbook_date_sheets(url_args)
    transaction_result('cleanup', '', backup_name, removed=True, sheetOrder=sheet_order)


def cleanup_temporary_transactions(url_args):
    """Remove every historical AMDC rollback sheet and verify none remain."""
    prefix = '__amdc_rollback_'
    temporary_sheets = [sheet for sheet in transaction_sheets(url_args) if sheet_title(sheet).startswith(prefix)]
    removed = []
    for sheet in temporary_sheets:
        title = sheet_title(sheet)
        cli(['+sheet-delete', *url_args, '--sheet-id', sheet['sheet_id'], '--yes'])
        removed.append(title)
    remaining = [sheet_title(sheet) for sheet in transaction_sheets(url_args) if sheet_title(sheet).startswith(prefix)]
    if remaining:
        raise RuntimeError(f'temporary transaction sheets remain: {remaining}')
    sheet_order = sort_workbook_date_sheets(url_args)
    transaction_result('cleanup-temporary', '', '', removedSheets=removed, removedCount=len(removed), sheetOrder=sheet_order)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', required=False)
    parser.add_argument('--week', required=False, help='YYYY-MM-DD')
    parser.add_argument('--url', default=configured_feishu_url())
    parser.add_argument('--transaction-action', choices=('prepare', 'rollback', 'cleanup', 'cleanup-temporary'))
    parser.add_argument('--transaction-backup', default='')
    parser.add_argument('--transaction-existed', default='true')
    parser.add_argument('--transaction-index', type=int, default=-1)
    args = parser.parse_args()
    if not args.url:
        fail('Feishu sheet URL is not configured. Set integrations.feishuSheetUrl in amdc-config.json, AMDC_FEISHU_SHEET_URL, or --url.')
    if args.week and not re.fullmatch(r'\d{4}-\d{2}-\d{2}', args.week):
        fail('week must use YYYY-MM-DD')
    target_name = (args.week or '').replace('-', '')
    url_args = ['--url', args.url.split('?', 1)[0]]
    if args.transaction_action:
        if args.transaction_action != 'cleanup-temporary' and not args.transaction_backup:
            fail('--transaction-backup is required for transaction actions')
        if args.transaction_action != 'cleanup-temporary' and not args.week:
            fail('--week is required for transaction actions')
        try:
            if args.transaction_action == 'prepare':
                prepare_transaction(url_args, target_name, args.transaction_backup)
            elif args.transaction_action == 'rollback':
                rollback_transaction(
                    url_args, target_name, args.transaction_backup,
                    str(args.transaction_existed).lower() == 'true', args.transaction_index,
                )
            elif args.transaction_action == 'cleanup':
                cleanup_transaction(url_args, args.transaction_backup)
            else:
                cleanup_temporary_transactions(url_args)
            return
        except Exception as error:
            fail(str(error))
    if not args.week:
        fail('--week is required for sync actions')
    if not args.xlsx:
        fail('--xlsx is required for sync actions')
    xlsx = Path(args.xlsx).resolve()
    if not xlsx.is_file():
        fail(f'Excel file not found: {xlsx}')
    wb = load_workbook(xlsx, data_only=False)
    ws = wb.active
    max_row, max_col = ws.max_row, ws.max_column
    if max_row < 5 or max_col < 1:
        fail('Excel workbook has no report data')
    progress(12, '已读取 Excel，正在定位飞书工作表')

    try:
        info = cli(['+workbook-info', *url_args])['data']
        sheet = next((s for s in info.get('sheets', []) if (s.get('title') or s.get('sheet_name')) == target_name), None)
        created = False
        if not sheet:
            cli(['+sheet-create', *url_args, '--title', target_name, '--row-count', str(max(200, max_row)), '--col-count', str(max(20, max_col))])
            info = cli(['+workbook-info', *url_args])['data']
            sheet = next((s for s in info.get('sheets', []) if (s.get('title') or s.get('sheet_name')) == target_name), None)
            created = True
        if not sheet:
            raise RuntimeError(f'cannot resolve target sheet: {target_name}')
        sheet_id = sheet['sheet_id']
        existing_rows, existing_cols = int(sheet.get('row_count') or 0), int(sheet.get('column_count') or 0)
        existing_range = f'A1:{get_column_letter(max(existing_cols, max_col))}{max(existing_rows, max_row)}'

        progress(20, '正在检查应用标题和本周排名是否重复')
        if not created and report_columns_match(ws, max_row, url_args, sheet_id, existing_rows):
            progress(94, '正在按日期递减排列工作表')
            sheet_order = sort_workbook_date_sheets(url_args)
            progress(100, '标题和本周排名完全重复，已跳过同步')
            print(json.dumps({
                'ok': True, 'sheetId': sheet_id, 'sheetName': target_name, 'created': False,
                'skipped': True, 'skipReason': '标题列和本周排名列完全重复',
                'rows': max_row, 'columns': max_col, 'xlsx': str(xlsx), 'sheetOrder': sheet_order,
            }, ensure_ascii=True))
            return

        if not created and existing_rows:
            cli(['+cells-unmerge', *url_args, '--sheet-id', sheet_id, '--range', existing_range])
            cli(['+cells-clear', *url_args, '--sheet-id', sheet_id, '--range', existing_range, '--scope', 'all', '--yes'])
        progress(25, '已清理旧内容，正在写入数据')
        if existing_rows < max_row:
            cli(['+dim-insert', *url_args, '--sheet-id', sheet_id, '--position', str(max(0, existing_rows - 1)), '--count', str(max_row - existing_rows)])
        if existing_cols < max_col:
            cli(['+dim-insert', *url_args, '--sheet-id', sheet_id, '--position', get_column_letter(existing_cols + 1), '--count', str(max_col - existing_cols)])

        for start in range(1, max_row + 1, CHUNK_ROWS):
            end = min(max_row, start + CHUNK_ROWS - 1)
            cells = [[cell_payload(ws.cell(row, col), no_wrap=row >= 5) for col in range(1, max_col + 1)] for row in range(start, end + 1)]
            cli(['+cells-set', *url_args, '--sheet-id', sheet_id, '--range', f'A{start}:{get_column_letter(max_col)}{end}', '--cells', '-'], input_data=json.dumps(cells, ensure_ascii=False))
            progress(25 + round(end * 50 / max_row), f'正在写入数据：{end}/{max_row} 行')

        operations = []
        for merged in ws.merged_cells.ranges:
            operations.append({'shortcut': '+cells-merge', 'input': {'sheet_id': sheet_id, 'range': str(merged)}})
        for col in range(1, max_col + 1):
            pixels = column_pixels(ws, col)
            if pixels:
                letter = get_column_letter(col)
                operations.append({'shortcut': '+cols-resize', 'input': {'sheet_id': sheet_id, 'range': letter, 'type': 'pixel', 'size': pixels}})
        for row in range(1, max_row + 1):
            height = ws.row_dimensions[row].height
            if height:
                operations.append({'shortcut': '+rows-resize', 'input': {'sheet_id': sheet_id, 'range': str(row), 'type': 'pixel', 'size': max(16, round(height * 4 / 3))}})
        if operations:
            progress(80, '正在按最长内容调整列宽与版式')
            cli(['+batch-update', *url_args, '--operations', '-', '--yes'], input_data=json.dumps(operations, ensure_ascii=False))

        pane = ws.freeze_panes
        if pane:
            progress(90, '正在还原冻结窗格')
            row_count = max(0, ws[pane].row - 1)
            col_count = max(0, ws[pane].column - 1)
            cli(['+dim-freeze', *url_args, '--sheet-id', sheet_id, '--dimension', 'row', '--count', str(row_count)])
            cli(['+dim-freeze', *url_args, '--sheet-id', sheet_id, '--dimension', 'column', '--count', str(col_count)])

        cli(['+filter-delete', *url_args, '--sheet-id', sheet_id, '--yes'], allow_failure=True)

        progress(94, '正在按日期递减排列工作表')
        sheet_order = sort_workbook_date_sheets(url_args)

        progress(96, '正在回读校验同步结果')
        check = cli(['+csv-get', *url_args, '--sheet-id', sheet_id, '--range', f'A1:{get_column_letter(max_col)}5', '--max-chars', '12000'])['data']
        content = check.get('annotated_csv', '')
        if target_name not in content or '序号' not in content:
            raise RuntimeError('Feishu read-back validation failed: report title or header is missing')
        progress(100, '同步完成')
        print(json.dumps({
            'ok': True, 'sheetId': sheet_id, 'sheetName': target_name, 'created': created,
            'rows': max_row, 'columns': max_col, 'xlsx': str(xlsx), 'sheetOrder': sheet_order,
        }, ensure_ascii=True))
    except Exception as error:
        fail(str(error))


if __name__ == '__main__':
    main()
