# -*- coding: utf-8 -*-
import datetime
import re
import unicodedata
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def previous_monday():
    d = datetime.date.today()
    return (d - datetime.timedelta(days=d.weekday() + 7)).isoformat()


HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(bold=True, size=14, color='1F3864')
NOTE_FONT = Font(size=9, color='595959')
LINK_FONT = Font(color='0563C1', underline='single', size=10)
LINK_FONT_BOLD = Font(color='0563C1', underline='single', size=10, bold=True)
UP_FILL = PatternFill('solid', fgColor='E2EFDA')
DOWN_FILL = PatternFill('solid', fgColor='FCE4E4')
NEW_FILL = PatternFill('solid', fgColor='FFF2CC')
FOCUS_FILL = PatternFill('solid', fgColor='F8CBAD')
SUSPECTED_DELISTED_TITLE_FILL = PatternFill('solid', fgColor='FFC7CE')
POTENTIAL_NEW_TITLE_FILL = PatternFill('solid', fgColor='C6E0B4')
thin = Side(style='thin', color='D9D9D9')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
GROUP_TOP = Border(left=thin, right=thin, bottom=thin, top=Side(style='medium', color='1F3864'))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=False)
LEFT = Alignment(horizontal='center', vertical='center', wrap_text=False)

COLS = [
    ('序号', 5, 'c'),
    ('应用标题', 28, 'l'),
    ('品类', 12, 'c'),
    ('Tag路径', 22, 'l'),
    ('本周排名', 7, 'c'),
    ('上周排名', 7, 'c'),
    ('变化量', 7, 'c'),
    ('6周排名轨迹', 30, 'l'),
    ('Top50稳定性', 20, 'l'),
    ('近30天下载国Top5', 30, 'l'),
    ('近30天收入国Top5', 26, 'l'),
    ('市场属性', 20, 'l'),
    ('上线日期', 12, 'c'),
    ('评分', 6, 'c'),
    ('评论数', 11, 'c'),
    ('发行商', 20, 'l'),
    ('总部', 6, 'l'),
    ('重点关注', 8, 'c'),
    ('备注', 26, 'l'),
]


def build_column_index(cols=COLS):
    return {name: i + 1 for i, (name, _, _) in enumerate(cols)}


def visual_width(value):
    text = str(value or '').replace('\r', ' ').replace('\n', ' ')
    return sum(2 if unicodedata.east_asian_width(ch) in ('W', 'F', 'A') else 1 for ch in text)


def finalize_worksheet(ws, head_row, cols=COLS):
    """Keep every populated cell on one centered line and size columns to content."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, str):
                cell.value = cell.value.replace('\r', ' ').replace('\n', ' ')
            cell.alignment = CENTER

    for row_idx in range(head_row, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20

    for col_idx, (_, preset_width, _) in enumerate(cols, 1):
        content_width = max(
            (visual_width(ws.cell(row_idx, col_idx).value) for row_idx in range(head_row, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(preset_width, content_width + 2))

    for row_idx in (1, 2, 3):
        if ws.cell(row_idx, 1).value is None:
            continue
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(cols))
        ws.cell(row_idx, 1).alignment = CENTER
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 18


def chain_taxonomy(tax):
    byid = {t['id']: t for t in tax}
    parents = set(p for t in tax for p in (t.get('parent_ids') or []))
    leaves = [t for t in tax if t['id'] not in parents] or tax
    best = []
    for lf in leaves:
        chain, cur, seen = [], lf, set()
        while cur and cur['id'] not in seen:
            seen.add(cur['id'])
            chain.append(cur['name'])
            pid = (cur.get('parent_ids') or [None])[0]
            cur = byid.get(pid)
        chain = list(reversed(chain))
        if len(chain) > len(best):
            best = chain
    return best


def tag_path(tags):
    tags = tags or []
    domain = next((t['name'] for t in tags if t.get('type') == 'domain'), '')
    games = [t for t in tags if t.get('type') == 'games']
    seq = [domain] if domain else []
    if games:
        meta = next((t['name'] for t in tags if t.get('type') == 'meta'), '')
        if meta:
            seq.append(meta)
        seq += chain_taxonomy(games)
    else:
        seq += chain_taxonomy([t for t in tags if t.get('type') == 'apps'])
    out = []
    for x in seq:
        if x and x not in out:
            out.append(x)
    return ' / '.join(out[:3])


def norm_date(s):
    if not s:
        return ''
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    return f'{m.group(1)}-{m.group(2)}-{m.group(3)}' if m else s


def hist_str(history):
    return ' → '.join(str(x) if x is not None else '·' for x in reversed(history))


def stability(r):
    history = r['history']
    count50 = sum(1 for h in history if h is not None and h <= 50)
    onboard = sum(1 for h in history if h is not None)
    parts = []
    if onboard > 0:
        parts.append(f'在榜{onboard}周')
    if count50 > 0:
        parts.append(f'Top50 {count50}周')
    return '/'.join(parts)


def change_str(r):
    if r['lastWeek'] is None:
        return 'NEW'
    c = r['change']
    return f'+{c}' if c > 0 else (str(c) if c < 0 else '0')


def top5(s):
    if not s:
        return ''
    return ' / '.join([x for x in s.split(' / ') if x][:5])


def store_url(store_ids):
    sids = store_ids or []
    gp = next((s for s in sids if s.startswith('1_')), None)
    if gp:
        return 'https://play.google.com/store/apps/details?id=' + gp[2:]
    ios = next((s for s in sids if s.startswith('2_')), None) or next((s for s in sids if s.startswith('3_')), None)
    if ios:
        return 'https://apps.apple.com/app/id' + ios[2:]
    return ''


def suspected_delisted(r):
    # 国别为空本身不再等同于下架；采集端会先核验商店链接，只有明确
    # 返回不存在并写入“默认下架”时才在导出中标记。
    return r.get('countryStatus') == '默认下架'


def potential_new(r):
    return bool(r.get('potentialNew')) or any('潜力新品' in str(reason) for reason in (r.get('_focusReasons') or []))


def build_focus_cells(r, cat):
    co = r.get('country') or {}
    reasons = '/'.join(r.get('_focusReasons', []))
    app_name = r.get('name', '')
    if app_name and not app_name.endswith('（疑似下架）') and suspected_delisted(r):
        app_name = f'{app_name}（疑似下架）'
    return {
        '序号': None,
        '应用标题': app_name,
        '品类': cat,
        'Tag路径': tag_path(r.get('tags')),
        '本周排名': r['rank'],
        '上周排名': r['lastWeek'] if r['lastWeek'] is not None else 'NEW',
        '变化量': change_str(r),
        '6周排名轨迹': hist_str(r['history']),
        'Top50稳定性': stability(r),
        '近30天下载国Top5': top5(co.get('dlList', '')) or '—',
        '近30天收入国Top5': top5(co.get('revList', '')) or '—（无内购收入/IAA变现）',
        '市场属性': co.get('market', ''),
        '上线日期': norm_date(r.get('release', '')),
        '评分': round(r['rating'], 2) if r.get('rating') else '—',
        '评论数': f"{r['reviews']:,}" if r.get('reviews') else '—',
        '发行商': r.get('publisher', ''),
        '总部': r.get('hq', ''),
        '重点关注': '是' if r.get('_focus', False) else '',
        '备注': reasons if reasons else '',
    }


def apply_focus_row(ws, row, seq, r, cat, col_map, cols=COLS, *, first_in_group=False, category_fill=None):
    cells = build_focus_cells(r, cat)
    cells['序号'] = seq
    border = GROUP_TOP if first_in_group else BORDER
    for name, _, align in cols:
        j = col_map[name]
        cell = ws.cell(row, j, cells[name])
        cell.border = border
        cell.alignment = CENTER if align == 'c' else LEFT

    name_cell = ws.cell(row, col_map['应用标题'])
    url = r.get('storeLink') or store_url(r.get('storeIds'))
    if url:
      name_cell.hyperlink = url
      name_cell.font = LINK_FONT_BOLD if r.get('_focus', False) else LINK_FONT

    change_cell = ws.cell(row, col_map['变化量'])
    if r['lastWeek'] is None:
        change_cell.fill = NEW_FILL
    elif r['change'] is not None and r['change'] > 0:
        change_cell.fill = UP_FILL
    elif r['change'] is not None and r['change'] < 0:
        change_cell.fill = DOWN_FILL

    if category_fill is not None:
        ws.cell(row, col_map['品类']).fill = category_fill

    if suspected_delisted(r):
        name_cell.fill = SUSPECTED_DELISTED_TITLE_FILL
    elif potential_new(r):
        name_cell.fill = POTENTIAL_NEW_TITLE_FILL
    elif r.get('_focus', False):
        name_cell.fill = FOCUS_FILL
