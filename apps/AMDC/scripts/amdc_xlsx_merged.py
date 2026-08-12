# -*- coding: utf-8 -*-
"""Merge category focus rows into one workbook."""
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from amdc_xlsx_common import (
    COLS,
    HEAD_FILL,
    HEAD_FONT,
    TITLE_FONT,
    NOTE_FONT,
    build_column_index,
    apply_focus_row,
    previous_monday,
    finalize_worksheet,
)


PROJECT_DIR = Path(os.environ.get('AMDC_PROJECT_DIR', Path.cwd())).resolve()


def u(*codes):
    return ''.join(chr(x) for x in codes)


ORDER = [
    u(0x8D85, 0x4F11, 0x95F2),
    u(0x4F11, 0x95F2),
    u(0x58C1, 0x7EB8),
    'Launcher',
    u(0x6740, 0x6BD2, 0x8F6F, 0x4EF6, 0x3001, 0x6E05, 0x7406),
    u(0x6587, 0x4EF6, 0x6062, 0x590D),
    'PDF' + u(0x9605, 0x8BFB, 0x5668),
]
CAT_FILLS = {
    ORDER[0]: 'DDEBF7',
    ORDER[1]: 'E2EFDA',
    ORDER[2]: 'D9EAD3',
    'Launcher': 'FCE4D6',
    ORDER[3]: 'FFF2CC',
    ORDER[4]: 'EDEDED',
    ORDER[5]: 'E1D5E7',
}


def load(cat, data_dir):
    p = data_dir / f'amdc-{cat}-weekly.json'
    if not os.path.exists(p):
        print('  [skip] no data:', cat)
        return None
    d = json.load(open(p, encoding='utf-8'))
    d['_mon'] = (d.get('weeks') or [''])[0].replace('-', '')
    return d


def normalize_top_depth(data):
    inferred = 100 if data.get('topDepth') is None and len(data.get('records') or []) <= 100 else 1000
    value = data.get('topDepth') or os.environ.get('TOP_DEPTH') or inferred
    try:
        depth = int(value)
    except (TypeError, ValueError):
        depth = 100
    return 100 if depth == 100 else 1000


def main():
    anchor = os.environ.get('WEEK_ANCHOR') or previous_monday()
    mon = anchor.replace('-', '')
    cache_dir = Path(os.environ.get('AMDC_RUN_DIR', PROJECT_DIR / 'Cache' / mon)).resolve()

    groups = []
    mdef = None
    gen = ''
    anchors = {}
    depths = {}
    for cat in ORDER:
        d = load(cat, cache_dir)
        if not d:
            continue
        mdef = mdef or d.get('marketDef')
        gen = gen or d.get('generatedAt', '')[:10]
        anchors[cat] = (d.get('weeks') or [''])[0]
        depths[cat] = normalize_top_depth(d)
        groups.append((cat, d['_mon'], sorted(d['focus'], key=lambda r: r['rank'])))

    distinct = sorted(set(a for a in anchors.values() if a))
    if len(distinct) > 1 and os.environ.get('ALLOW_MIXED_WEEKS') != '1':
        detail = ' / '.join(f'{c}={a}' for c, a in anchors.items())
        raise SystemExit(
            f'[abort] 合并的品类周锚点不一致: {detail}\n'
            '请用同一 WEEK_ANCHOR 重跑落后的品类，或设 ALLOW_MIXED_WEEKS=1 强制合并（标题会标注混周）。'
        )
    mixed = len(distinct) > 1
    mon = max((_mon for _, _mon, _ in groups), default='')
    col_map = build_column_index(COLS)

    wb = Workbook()
    ws = wb.active
    ws.title = mon if mon else 'Sheet1'

    mix_note = f'  ⚠ 混周合并：{" / ".join(f"{c}={a}" for c, a in anchors.items())}' if mixed else ''
    distinct_depths = sorted(set(depths.values())) or [100]
    depth_label = 'Top' + (str(distinct_depths[0]) if len(distinct_depths) == 1 else '/'.join(str(v) for v in distinct_depths))
    ws['A1'] = f'AMDC 周报 · 全品类（{"/".join(c for c, _, _ in groups)}）· {mon} 当周（免费榜）{mix_note}'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = (
        f'口径：全球(WW)·周聚合·免费榜·{depth_label} | '
        '变化量=正数上升/负数下降/NEW首次出现 | '
        '重点关注=首次进入Top100或变化突出(前10绝对↑≥5 / 10-200相对↑>50%)；潜力新品=首进前100+最多3周排名记录+陌生发行商+成熟市场下载占比或收入占比任一≥25% | '
        '数据源 AMDC API · 生成 ' + gen
    )
    ws['A2'].font = NOTE_FONT
    if mdef:
        ws['A3'] = '市场定义：成熟市场(高ARPU)= ' + ' '.join(mdef['mature']) + '  ｜  新兴市场= ' + ' '.join(mdef['emerging'])
        ws['A3'].font = NOTE_FONT

    head_row = 5
    for name, width, _ in COLS:
        j = col_map[name]
        cell = ws.cell(head_row, j, name)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        ws.column_dimensions[get_column_letter(j)].width = width

    row = head_row + 1
    seq = 0
    for cat, _mon, focus in groups:
        first = True
        category_fill = PatternFill('solid', fgColor=CAT_FILLS.get(cat, 'FFFFFF'))
        for record in focus:
            seq += 1
            apply_focus_row(
                ws,
                row,
                seq,
                record,
                cat,
                col_map,
                first_in_group=first,
                category_fill=category_fill,
            )
            first = False
            row += 1

    ws.freeze_panes = ws.cell(head_row + 1, 4)
    ws.auto_filter.ref = f'A{head_row}:{get_column_letter(len(COLS))}{row - 1}'
    finalize_worksheet(ws, head_row, COLS)
    output_dir = PROJECT_DIR / 'output'
    output_dir.mkdir(parents=True, exist_ok=True)
    out = Path(os.environ.get('AMDC_MERGED_XLSX') or (output_dir / f'AMDC-{mon}.xlsx'))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print('saved', out, '| 品类', len(groups), '| 总行', row - head_row - 1)


if __name__ == '__main__':
    main()
