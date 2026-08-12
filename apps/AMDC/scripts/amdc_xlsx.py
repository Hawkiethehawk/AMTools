# -*- coding: utf-8 -*-
"""Read amdc-<cat>-weekly.json and export one category workbook."""
import json
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from amdc_xlsx_common import (
    COLS,
    HEAD_FILL,
    HEAD_FONT,
    TITLE_FONT,
    NOTE_FONT,
    build_column_index,
    apply_focus_row,
    previous_monday,
    finalize_worksheet,
)


PROJECT_DIR = Path(os.environ.get('AMDC_PROJECT_DIR', Path.cwd())).resolve()


def top_depth_label(data):
    inferred = 100 if data.get('topDepth') is None and len(data.get('records') or []) <= 100 else 1000
    value = data.get('topDepth') or os.environ.get('TOP_DEPTH') or inferred
    try:
        depth = int(value)
    except (TypeError, ValueError):
        depth = 100
    return f'Top{100 if depth == 100 else 1000}'


def main():
    anchor = os.environ.get('WEEK_ANCHOR') or previous_monday()
    mon = anchor.replace('-', '')
    out_base = Path(os.environ.get('AMDC_RUN_DIR', PROJECT_DIR / 'Cache' / mon)).resolve()
    default_cat = ''.join(chr(x) for x in [0x8D85, 0x4F11, 0x95F2])
    cat = sys.argv[1] if len(sys.argv) > 1 else default_cat
    src = out_base / f'amdc-{cat}-weekly.json'

    data = json.load(open(src, encoding='utf-8'))
    weeks = data['weeks']
    mon = (weeks[0] or '').replace('-', '') or mon
    out = out_base / f'AMDC-{cat}-{mon}.xlsx'
    focus = sorted(data['focus'], key=lambda r: r['rank'])
    mdef = data.get('marketDef', {'mature': [], 'emerging': []})
    col_map = build_column_index(COLS)

    wb = Workbook()
    ws = wb.active
    ws.title = f'{cat}-重点{len(focus)}'

    ws['A1'] = f'AMDC 周报 · {cat} · {weeks[0]} 当周（免费榜）'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = (
        f'口径：全球(WW)·周聚合·免费榜·{top_depth_label(data)} | '
        '变化量=正数上升/负数下降/NEW首次出现 | '
        '重点关注=首次进入Top100或变化突出(前10绝对↑≥5 / 10-200相对↑>50%)；潜力新品=首进前100+最多3周排名记录+陌生发行商+成熟市场下载占比或收入占比任一≥25% | '
        '数据源 AMDC API · 生成 ' + data['generatedAt'][:10]
    )
    ws['A2'].font = NOTE_FONT
    ws['A3'] = '市场定义：成熟市场(高ARPU)= ' + ' '.join(mdef['mature']) + '   ｜   新兴市场= ' + ' '.join(mdef['emerging'])
    ws['A3'].font = NOTE_FONT

    head_row = 5
    for name, width, _ in COLS:
        j = col_map[name]
        cell = ws.cell(head_row, j, name)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        ws.column_dimensions[get_column_letter(j)].width = width

    row = head_row + 1
    for idx, record in enumerate(focus, 1):
        apply_focus_row(ws, row, idx, record, cat, col_map)
        row += 1

    ws.freeze_panes = ws.cell(head_row + 1, 3)
    ws.auto_filter.ref = f'A{head_row}:{get_column_letter(len(COLS))}{row - 1}'
    finalize_worksheet(ws, head_row, COLS)
    out_base.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print('saved', out, '| focus', len(focus))


if __name__ == '__main__':
    main()
